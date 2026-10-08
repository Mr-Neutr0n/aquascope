"""Tests for the Studio deliverables: figures, tables, workbook, reports, notebook and bundle.

The payloads are synthetic and given (nothing is computed), in the shapes
the tools return; the workspace is built by hand.
"""

from __future__ import annotations

import io
import json
import math
import random
import re
import sys
import zipfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import pytest  # noqa: E402

from aquascope.studio import catalogue  # noqa: E402
from aquascope.studio.deliverables import (  # noqa: E402
    _common,
    bundle,
    figures,
    notebook,
    report_docx,
    report_md,
    tables,
    workbook,
)
from aquascope.studio.workspace import Brief, Dataset, Inventory, Workspace  # noqa: E402

PNG = b"\x89PNG\r\n\x1a\n"
SITE = {"lat": 51.4, "lon": -0.3}
RP = [2, 5, 10, 25, 50, 100]


# ── synthetic payloads ────────────────────────────────────────────────────


def _daily(years: int = 30, start: int = 1990) -> tuple[list[str], list[float]]:
    import datetime as dt

    rng = random.Random(1)
    t, v = [], []
    d, end = dt.date(start, 1, 1), dt.date(start + years, 1, 1)
    while d < end:
        season = 15 * math.sin(2 * math.pi * d.timetuple().tm_yday / 365.25)
        v.append(round(20 + season + rng.gauss(0, 3), 2))
        t.append(d.isoformat())
        d += dt.timedelta(days=1)
    return t, v


def _months(n: int = 240, start: int = 2005) -> list[str]:
    out, y, m = [], start, 1
    for _ in range(n):
        out.append(f"{y:04d}-{m:02d}-01")
        m += 1
        if m > 12:
            m, y = 1, y + 1
    return out


@pytest.fixture(scope="module")
def station() -> dict:
    t, v = _daily()
    years = sorted({int(x[:4]) for x in t})
    am = [max(val for d, val in zip(t, v) if int(d[:4]) == y) for y in years]
    gev = [40, 46, 50, 55, 58, 62]
    return {
        "source": "uk_ea", "station_id": "3400TH", "variable": "discharge", "unit": "m3/s", "n": len(t),
        "start": t[0], "end": t[-1], "years": 30.0,
        "stats": {"mean": 20.1, "median": 19.8, "min": 2.1, "max": 66.0},
        "series": {"t": t, "v": v},
        "annual_max": {"year": years, "v": am},
        "fdc": {"exceedance": list(range(1, 100)), "q": [60 * math.exp(-0.035 * i) for i in range(1, 100)],
                "q95": 3.0, "q50": 18.0, "q10": 35.0},
        "ffa": {"n_years": 30, "return_periods": RP, "fits": {
            "gev_lmoments": {"q": gev, "params": [0.1, 40, 5]},
            "lp3": {"q": [39, 45, 49, 54, 57, 60],
                    "ci": [[36, 42], [42, 48], [45, 53], [49, 60], [51, 64], [53, 69]], "params": [1, 2, 3]},
            "gev_bootstrap": {"q": gev, "ci": [[37, 43], [43, 49], [46, 54], [50, 61], [52, 66], [54, 72]],
                              "params": [0.1, 40, 5], "n_bootstrap": 50, "n_bootstrap_discarded": 0},
        }},
        "trend": {"on": "annual mean", "p_value": 0.21, "tau": 0.12, "trend": "no trend",
                  "sens_slope_per_year": 0.05, "n_years": 30},
        "methods": [], "notes": [],
    }


@pytest.fixture(scope="module")
def drought() -> dict:
    rng = random.Random(2)
    mi = _months()

    def idx() -> list[float]:
        return [round(rng.gauss(0, 1), 3) for _ in mi]

    rows = []
    for s in (1, 3, 12):
        rows.append({"timescale": s,
                     "spi": {"current": -0.4, "class": "near_normal", "date": mi[-1], "worst": -2.5,
                             "worst_date": "2012-08-01", "n": len(mi), "events": 4},
                     "spei": {"current": -0.9, "class": "near_normal", "date": mi[-1], "worst": -2.8,
                              "worst_date": "2012-08-01", "n": len(mi), "events": 5},
                     "divergence": {"current": -0.5, "mean_last_10y": -0.2, "months_spei_drier_pct": 60.0,
                                    "correlation": 0.9, "n": len(mi)},
                     "series": {"index": mi, "step": 1, "spi": idx(), "spei": idx()}})
    return {"latitude": 51.4, "longitude": -0.3, "timescales": [1, 3, 12], "headline_timescale": 3,
            "headline_index": "spei", "threshold": -1.0, "months": len(mi), "start": mi[0], "end": mi[-1],
            "years": 20.0, "indices": rows, "status": "near_normal", "in_drought": False,
            "temperature": {"mean_c": 10.2, "trend_c_per_decade": 0.3, "p_value": 0.01, "trend": "increasing",
                            "n_years": 20}, "methods": [], "notes": []}


@pytest.fixture(scope="module")
def donors() -> dict:
    rng = random.Random(3)
    return {"latitude": 51.4, "longitude": -0.3, "method": "combined", "k": 5, "n_candidates": 900,
            "stations": [{"source": "uk_ea", "station_id": f"S{i}", "name": f"Donor {i}",
                          "latitude": 51.4 + rng.uniform(-1, 1), "longitude": -0.3 + rng.uniform(-1.5, 1.5),
                          "score": round(rng.random(), 3), "distance_km": round(rng.uniform(10, 150), 1),
                          "up_area_km2": 120.0 + i} for i in range(5)]}


@pytest.fixture(scope="module")
def climate() -> dict:
    return {"latitude": 51.4, "longitude": -0.3, "start": "2015-01-01", "end": "2025-01-01", "years": 10,
            "climate": {"source": "ERA5 via Open-Meteo", "precipitation_mm_per_year": 650.0,
                        "et0_mm_per_year": 580.0, "temperature_mean_c": 10.5, "aridity_index": 1.1,
                        "aridity_class": "humid",
                        "monthly_precipitation_mm": [60, 45, 40, 42, 48, 50, 45, 55, 50, 70, 68, 62],
                        "monthly_et0_mm": [10, 18, 35, 60, 85, 95, 100, 85, 55, 28, 12, 8]},
            "glofas": {"variable": "discharge", "unit": "m3/s", "n": 7000, "start": "2005-01-01",
                       "end": "2025-01-01", "years": 20.0, "stats": {"mean": 12.0},
                       "annual_max": {"year": list(range(2005, 2025)), "v": [30 + (i % 7) for i in range(20)]},
                       "fdc": {"q95": 2.0, "q50": 10.0, "q10": 25.0}, "modelled": True}}


@pytest.fixture(scope="module")
def samples() -> dict:
    rng = random.Random(4)
    rows = [{"datetime": f"202{i % 5}-0{1 + i % 9}-15T00:00:00", "parameter": p,
             "value": round(base + rng.gauss(0, sd), 2), "unit": u}
            for i in range(20) for p, base, sd, u in (("pH", 7.4, 0.3, "std units"),
                                                       ("dissolved_oxygen", 9.0, 1.5, "mg/L"))]
    return {"source": "usgs", "station_id": "01646500", "n_samples": len(rows), "start": "2020-01-01",
            "end": "2024-12-01", "samples": rows,
            "parameters": {"pH": {"n": 20, "unit": "std units", "start": "2020-01-15", "end": "2024-09-15",
                                  "min": 6.8, "median": 7.4, "max": 8.0}},
            "sample_counts": {"pH": 20, "dissolved_oxygen": 20}, "units": {"pH": "std units"}}


@pytest.fixture(scope="module")
def others(station: dict, donors: dict) -> dict[str, dict]:
    """The smaller payloads, by the figure kind they carry."""
    t, v = station["series"]["t"][:1500], station["series"]["v"][:1500]
    mi = _months(120)
    rng = random.Random(5)
    sgi = [round(rng.gauss(0, 1), 3) for _ in mi]
    return {
        "propagation": {"source": "bgs", "station_id": "W1", "latitude": 51.4, "longitude": -0.3,
                        "sgi": {"threshold": -1.0, "n": 120},
                        "propagation": {"best": {"timescale": 12, "lag_months": 3, "correlation": 0.71, "n": 110},
                                        "by_timescale": {"1": {"lag_months": 1, "correlation": 0.2, "n": 110},
                                                         "12": {"lag_months": 3, "correlation": 0.71, "n": 110}},
                                        "max_lag_months": 24},
                        "series": {"index": mi, "step": 1, "sgi": sgi, "spi": sgi[::-1]}},
        "signatures_band": {"latitude": 51.4, "longitude": -0.3, "method": "similarity",
                            "estimates": {"q_mean_mm": {"value": 1.2, "unit": "mm/d", "label": "Mean daily flow",
                                                        "low": 0.8, "high": 1.7, "n_donors": 5},
                                          "bfi": {"value": 0.55, "unit": "fraction", "label": "Baseflow index",
                                                  "low": 0.4, "high": 0.7, "n_donors": 5}},
                            "similarity": {"donors": donors["stations"], "k": 5},
                            "skill": {"by_signature": {"q_mean_mm": {"n": 300, "nse": 0.62}}}},
        "reliability_curve": {"source": "uk_ea", "station_id": "3400TH", "unit": "m3/s", "mode": "gauged",
                              "fdc": {"q05": 45.0, "q50": 18.0, "q90": 4.5, "q95": 3.0}, "reserve_m3s": 3.0,
                              "required_flow_m3s": 8.0, "reliability": {"daily": 0.82, "annual": 0.4},
                              "verdict": "marginal"},
        "demand_monthly": {"latitude": 51.4, "longitude": -0.3, "crop": "maize", "area_ha": 10.0,
                           "planting_month": 4, "season_days": 150,
                           "demand": {"gross_irrigation_mm": 343.0},
                           "per_season": [{"year": y, "etc_mm": 400 + i, "effective_rain_mm": 170 + i,
                                           "net_irrigation_mm": 230 + i, "gross_irrigation_mm": 330 + i}
                                          for i, y in enumerate(range(2015, 2020))]},
        "et0_monthly": {"eto": {"index": t, "values": [max(0.2, 3 + 2.5 * math.sin(2 * math.pi * i / 365.25))
                                                       for i in range(len(t))], "n": len(t), "step": 1},
                        "mean_mm_per_day": 3.0, "total_mm": 4500.0},
        "site_map": {"latitude": 51.4, "longitude": -0.3, "sub_basin": {"hybas_id": 2120000010},
                     "attributes": {"area_km2": 9900.0,
                                    "elevation_m": {"value": 120.0, "unit": "m", "label": "Mean elevation",
                                                    "source": "area_weighted_mean"}}},
        "who_exceedances": {"rows": [{"parameter": "nitrate", "rule": "at most 50 mg/L", "n": 20, "n_exceed": 4,
                                      "pct": 20.0, "status": "Alert"},
                                     {"parameter": "arsenic", "rule": "at most 0.01 mg/L", "n": 20,
                                      "n_exceed": 0, "pct": 0.0, "status": "OK"}], "n_alerts": 1},
        "wqi_bars": {"use": "drinking", "guideline_set": "drinking", "index": "ccme_wqi", "score": 78.3,
                     "category": "Fair", "n_samples": 40,
                     "ccme": {"score": 78.3, "category": "Fair", "f1": 20.0, "f2": 10.0, "f3": 12.5},
                     "nsf": {"score": 70.1, "category": "Medium"}},
        "baseflow": {"column": "discharge", "method": "lyne_hollick", "bfi": 0.61,
                     "series": {"index": t, "total": v, "baseflow": [x * 0.6 for x in v]}},
        "recharge": {"column": "level", "method": "water_table_fluctuation", "value_mm_per_year": 210.0,
                     "uncertainty": 60.0, "metadata": {"specific_yield": 0.15, "total_rise_m": 42.0}},
        "sgi_drought": {"column": "level", "sgi": {"index": mi, "values": sgi}, "threshold": -1.0,
                        "events": [{"start": "2008-05-01", "end": "2009-02-01", "duration": 10,
                                    "severity": -12.0, "peak": -2.9}]},
    }


@pytest.fixture(scope="module")
def recon() -> dict:
    return {"point": SITE, "stations": [
        {"source": "uk_ea", "station_id": "3400TH", "name": "Kingston", "latitude": 51.41, "longitude": -0.31,
         "distance_km": 1.2, "variables": ["discharge"], "period_start": "1990-01-01", "years": 35.0},
        {"source": "bgs", "station_id": "W1", "name": "Well", "latitude": 51.5, "longitude": -0.2,
         "distance_km": 12.0, "variables": ["groundwater_level"]}],
        "catchment": {"area_km2": 9900.0},
        "context": {"years_by_variable": {"discharge": 35.0}, "area_km2": 9900.0, "donors": 10, "ungauged": False},
        "sufficiency": [{"method": "at_site_flood_frequency", "status": "defensible", "reason": "35 years",
                         "tool": "analyze_station", "station": {"source": "uk_ea", "station_id": "3400TH"}},
                        {"method": "spi", "status": "marginal", "reason": "ERA5 only", "tool": "drought_indices",
                         "station": None}], "notes": []}


@pytest.fixture(scope="module")
def ws(station: dict, drought: dict, recon: dict) -> Workspace:
    """A workspace after the Analysts and the Author: a brief, an inventory, a v3 study, results, artifacts, a
    report with every section."""
    w = Workspace(site=dict(SITE))
    w.brief = Brief(problem="Design flow for a road crossing at Kingston, 100-year", decision="size a culvert",
                    quantities=["the 100-year flow with a band"], kind="flood_risk", ready=True)
    w.inventory = Inventory(site=dict(SITE), datasets=[
        Dataset(id="uk_ea:3400TH", kind="station", variable="discharge", source="uk_ea", station_id="3400TH",
                name="Kingston", lat=51.41, lon=-0.31, distance_km=1.2, start="1990-01-01", end="2019-12-31",
                years=30.0, resolution="daily"),
        Dataset(id="era5", kind="reanalysis", variable="precipitation", source="openmeteo", lat=51.4, lon=-0.3),
    ], recon=recon)
    w.study = {
        "version": 3, "title": "Design flow at Kingston", "question": w.brief.problem, "author": "methodologist",
        "aquascope_version": "0.14.0", "problem": {"kind": "flood_risk", "site": dict(SITE)},
        "plan": {"objective": "the 100-year flow", "methodology": ["1. reconnaissance", "2. at-site FFA"],
                 "assumptions": ["stationarity"]},
        "steps": [
            {"id": "s1", "tool": "assess_site", "arguments": {"lat": 51.4, "lon": -0.3},
             "rationale": "what exists here", "expects": [{"check": "not_empty", "path": "stations"}]},
            {"id": "s2", "tool": "analyze_station", "arguments": {"source": "uk_ea", "station_id": "3400TH"},
             "method": "at_site_flood_frequency", "rationale": "the gauge record",
             "expects": [{"check": "min_years", "path": "years", "value": 20}], "depends_on": ["s1"],
             "outputs": [{"kind": "figure", "id": "fig-s2-frequency_curve", "caption": "the curve"}]},
            {"id": "s3", "tool": "drought_indices", "arguments": {"lat": 51.4, "lon": -0.3},
             "method": "spei", "rationale": "the climate context"},
        ]}
    w.run = {"ok": True, "stopped_at": None, "stop_reason": None, "results": [
        {"id": "s1", "tool": "assess_site", "arguments": {}, "ok": True, "result": recon,
         "gates": [{"check": "not_empty", "passed": True, "detail": "2 stations", "path": "stations"}],
         "gates_passed": True},
        {"id": "s2", "tool": "analyze_station", "arguments": {}, "ok": True, "result": station,
         "gates": [{"check": "min_years", "passed": True, "detail": "30 >= 20", "path": "years"}],
         "gates_passed": True},
        {"id": "s3", "tool": "drought_indices", "arguments": {}, "ok": True, "result": drought, "gates": [],
         "gates_passed": True},
    ]}
    for r in w.run["results"]:
        site = {**SITE, "stations": recon["stations"]}
        for a in figures.figures_for(r["id"], r["tool"], r["result"], site=site):
            w.add_artifact(a)
        for a in tables.tables_for(r["id"], r["tool"], r["result"]):
            w.add_artifact(a)
    w.report = {
        "title": "Design flow at Kingston", "answer": "The 100-year flow is about 62 m3/s (54 to 72).",
        "key_numbers": [{"label": "Q100", "value": 62, "unit": "m3/s", "step": "s2"},
                        {"label": "Record length", "value": 30, "unit": "years", "step": "s2"}],
        "sections": [
            {"id": "summary", "title": "Summary", "text": "The **100-year** flow is 62 m3/s.", "figures": [],
             "tables": []},
            {"id": "problem", "title": "Problem and decision", "text": "Size a culvert.\n\n- one\n- two",
             "figures": [], "tables": []},
            {"id": "site_data", "title": "Site and data", "text": "Two stations.", "figures": [],
             "tables": ["tab-s1-stations_within_reach", "tab-s1-sufficiency"]},
            {"id": "methodology", "title": "Methodology", "text": "1. reconnaissance\n2. FFA with `gev_lmoments`",
             "figures": [], "tables": []},
            {"id": "results-s2", "title": "Results: the gauge", "text": "See the curve.",
             "figures": ["fig-s2-frequency_curve", "fig-s2-series"], "tables": ["tab-s2-return_levels"]},
            {"id": "results-s3", "title": "Results: drought", "text": "Wet lately.",
             "figures": ["fig-s3-drought_strip"], "tables": ["tab-s3-index_divergence"]},
            {"id": "limitations", "title": "Limitations", "text": "Stationarity assumed.", "figures": [],
             "tables": []},
            {"id": "recommendations", "title": "Recommendations", "text": "Design at 72 m3/s.", "figures": [],
             "tables": []},
            {"id": "appendix", "title": "Appendix", "text": "The study file follows.", "figures": [], "tables": []},
        ],
        "not_established": ["the effect of climate change on the 100-year flow"],
        "recommendations": ["design at the upper band"],
        "references": ["Hosking, J. R. M. (1990). L-moments. J. R. Stat. Soc. B 52, 105-124."],
        "caveats": ["A playbook caveat, printed verbatim."],
        "footer": "Produced by AquaScope Studio; 1,550 tokens.",
    }
    w.charge("consultant", 100, 50)
    w.charge("author", 1000, 400)
    return w


# ── figures ───────────────────────────────────────────────────────────────


def _payload_for(kind: str, station: dict, drought: dict, donors: dict, climate: dict, samples: dict,
                 others: dict, recon: dict) -> tuple[dict, dict | None]:
    if kind in ("series", "annual_maxima", "frequency_curve", "fdc", "trend"):
        return station, None
    if kind == "drought_strip":
        return drought, None
    if kind == "donors_map":
        return donors, None
    if kind in ("monthly_climate", "glofas_series"):
        return climate, None
    if kind == "samples_by_parameter":
        return samples, None
    if kind == "site_map":
        return others["site_map"], {**SITE, "stations": recon["stations"]}
    if kind in ADVANCED_KINDS:
        return ADVANCED[ADVANCED_KINDS[kind]], None
    return others[kind], None


#: Real payloads of the advanced study steps (USGS 01013500, trimmed), by tool.
ADVANCED = json.loads((Path(__file__).parents[1] / "fixtures" / "advanced_payloads.json").read_text("utf-8"))
ADVANCED_KINDS = {"change_points": "change_points", "nonstationary_levels": "nonstationary_flood",
                  "pot_frequency": "pot_flood", "model_fit": "catchment_model", "scenario_bars": "catchment_model",
                  "projection_spread": "climate_projection", "regional_growth": "regional_flood"}


def test_the_advanced_tables_build_from_real_payloads() -> None:
    from aquascope.studio.deliverables import tables

    for tool, payload in ADVANCED.items():
        names = catalogue.get(tool).tables
        arts = tables.tables_for("t", tool, payload)
        built = {a.meta["name"] for a in arts}
        assert set(names) - {"annual_maxima"} <= built, (tool, names, built)
        assert all(a.data.count(b"\n") >= 2 for a in arts), tool


def test_every_catalogue_kind_has_a_maker() -> None:
    assert set(catalogue.figure_kinds()) <= set(figures.kinds())


def test_every_catalogue_kind_draws_png_and_svg(station, drought, donors, climate, samples, others, recon) -> None:
    for kind in catalogue.figure_kinds():
        payload, site = _payload_for(kind, station, drought, donors, climate, samples, others, recon)
        arts = figures.figures_for("t", "x", payload, site=site, kinds=[kind])
        assert [a.id for a in arts] == [f"fig-t-{kind}", f"fig-t-{kind}-svg"], kind
        png, svg = arts
        assert png.data.startswith(PNG), kind
        assert svg.data.lstrip().startswith((b"<svg", b"<?xml")), kind
        assert png.name == f"figures/t_{kind}.png" and svg.name == f"figures/t_{kind}.svg"
        assert png.media_type == "image/png" and svg.media_type == "image/svg+xml"
        assert png.kind == svg.kind == "figure" and png.step == "t"
        assert png.meta["kind"] == kind and png.meta["tool"] == "x"
        assert png.caption and png.caption.endswith(".")


def test_figures_for_follows_the_catalogue(station) -> None:
    arts = figures.figures_for("s2", "analyze_station", station, site=SITE)
    kinds = [a.meta["kind"] for a in arts if a.media_type == "image/png"]
    assert kinds == catalogue.get("analyze_station").figures
    series = next(a for a in arts if a.id == "fig-s2-series")
    assert "UK EA 3400TH" in series.caption and "annual maxima" in series.caption
    curve = next(a for a in arts if a.id == "fig-s2-frequency_curve")
    assert "bootstrap" in curve.caption and "Log-Pearson" in curve.caption
    trend = next(a for a in arts if a.id == "fig-s2-trend")
    assert "no trend" in trend.caption


def test_missing_keys_yield_nothing_and_never_raise() -> None:
    odd = {"series": {"t": [], "v": []}, "ffa": {}, "trend": None, "indices": "no", "stations": [{}],
           "fdc": {"q95": None}, "climate": {"monthly_precipitation_mm": [1, 2]}, "samples": [], "rows": [],
           "estimates": {}, "eto": {"index": []}, "sgi": {}, "propagation": {}, "glofas": {}}
    for kind in figures.kinds():
        assert figures.figures_for("x", "analyze_station", {}, kinds=[kind]) == []
        assert figures.figures_for("x", "analyze_station", odd, kinds=[kind]) == []
        assert figures.figures_for("x", "analyze_station", {"error": "no data"}, kinds=[kind]) == []
    assert figures.figures_for("x", "no_such_tool", {"series": {"t": ["2020-01-01"], "v": [1.0]}}) == []
    for name in tables.names():
        assert tables.tables_for("x", "analyze_station", {}, names=[name]) == []
        assert tables.tables_for("x", "analyze_station", odd, names=[name]) == []


def test_draw_returns_a_figure(station) -> None:
    fig = figures.draw("annual_maxima", station)
    assert fig is not None and fig.axes
    assert figures.png_bytes(fig).startswith(PNG)
    figures.close(fig)
    assert figures.draw("annual_maxima", {}) is None


def test_figures_from_workbench_shapes(station) -> None:
    """The workbench spells the same things differently: dated maxima, a single distribution with bounds."""
    t = station["series"]["t"]
    payload = {"column": "flow", "distribution": "gev", "return_periods": RP, "return_levels": [40, 46, 50, 55, 58, 62],
               "lower_bound": [37, 43, 46, 50, 52, 54], "upper_bound": [43, 49, 54, 61, 66, 72],
               "confidence_level": 0.95, "n_years": 30,
               "empirical": {"return_period": [31 / (31 - r) for r in range(1, 31)], "value": list(range(30, 60))},
               "annual_max": {"index": t[::365], "values": [30 + i % 5 for i in range(len(t[::365]))]}}
    arts = figures.figures_for("w", "return_periods", payload)
    assert [a.id for a in arts] == ["fig-w-frequency_curve", "fig-w-frequency_curve-svg"]
    points = {"source": "uk_ea", "station_id": "1", "unit": "m", "variable": "water_level",
              "points": [[d, 1.0 + i % 3] for i, d in enumerate(t[::30])]}
    assert [a.meta["kind"] for a in figures.figures_for("g", "get_timeseries", points)] == ["series", "series"]


# ── tables ────────────────────────────────────────────────────────────────


def test_every_catalogue_table_has_a_maker() -> None:
    listed = {t for e in catalogue.entries().values() for t in e.tables}
    assert listed <= set(tables.names())


def test_tables_for_headers_match_meta(station, drought, donors, samples, recon, climate, others) -> None:
    cases = [("analyze_station", station), ("drought_indices", drought), ("similar_basins", donors),
             ("water_quality_samples", samples), ("assess_site", recon), ("anywhere", climate),
             ("regionalize_signatures", others["signatures_band"]), ("supply_reliability",
                                                                     others["reliability_curve"]),
             ("crop_water_demand", others["demand_monthly"]), ("baseflow", others["baseflow"]),
             ("sgi_drought", others["sgi_drought"]), ("drought_propagation", others["propagation"]),
             ("who_screen", others["who_exceedances"]), ("wqi", others["wqi_bars"]),
             ("describe_catchment", others["site_map"]), ("reference_et", others["et0_monthly"])]
    for tool, payload in cases:
        arts = tables.tables_for("t", tool, payload)
        assert arts, tool
        for a in arts:
            assert a.kind == "table" and a.media_type == "text/csv" and a.step == "t"
            assert a.id == f"tab-t-{a.meta['name']}" and a.name == f"tables/t_{a.meta['name']}.csv"
            df = tables.frame_of(a)
            assert list(df.columns) == a.meta["columns"], a.id
            assert len(df) == a.meta["rows"] > 0, a.id
            assert a.data.split(b"\n", 1)[0].decode() == ",".join(a.meta["columns"])


def test_return_levels_table_columns(station) -> None:
    art = tables.tables_for("s2", "analyze_station", station, names=["return_levels"])[0]
    df = tables.frame_of(art)
    assert list(df.columns) == ["T", "estimator", "unit", "estimate", "lower", "upper",
                                "interval_method", "confidence_level"]
    lm = df[df.estimator == "gev_lmoments"]
    lp3 = df[df.estimator == "lp3"]
    boot = df[df.estimator == "gev_bootstrap"]
    assert lm["T"].tolist() == RP and lm["lower"].isna().all() and lm["upper"].isna().all()
    assert lp3["upper"].iloc[-1] == 69 and boot["upper"].iloc[-1] == 72
    series = tables.tables_for("s2", "analyze_station", station, names=["series"])[0]
    assert series.meta["columns"] == ["datetime", "value"] and series.meta["rows"] == station["n"]


def test_indices_monthly_is_wide(drought) -> None:
    art = tables.tables_for("s3", "drought_indices", drought, names=["indices_monthly"])[0]
    assert art.meta["columns"] == ["date", "spi_1", "spei_1", "spi_3", "spei_3", "spi_12", "spei_12"]
    assert art.meta["rows"] == drought["months"]


# ── workbook ──────────────────────────────────────────────────────────────


def test_workbook_loads_with_expected_sheets(ws) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(workbook.workbook_bytes(ws)))
    names = wb.sheetnames
    for fixed in ("README", "Inventory", "Plan", "Gates", "Figures", "Ledger"):
        assert fixed in names
    for a in ws.artifacts_of("table"):
        assert a.name.rsplit("/", 1)[-1].removesuffix(".csv") in names
    assert all(len(n) <= 31 for n in names) and len(set(names)) == len(names)
    readme = wb["README"]
    cells = [str(c.value) for row in readme.iter_rows() for c in row if c.value is not None]
    assert "Design flow at Kingston" in cells and any("62 m3/s" in c for c in cells)
    assert any("Q100" in c for c in cells) and any("zenodo" in c for c in cells)
    plan = wb["Plan"]
    assert [c.value for c in plan[1]][:2] == ["id", "tool"] and plan.freeze_panes == "A2"
    assert plan["A2"].value == "s1" and plan["B3"].value == "analyze_station"
    assert plan["A1"].font.bold
    gates = wb["Gates"]
    assert gates.max_row == 3 and gates["B2"].value == "not_empty"
    ledger = wb["Ledger"]
    assert ledger["A2"].value == "consultant" and ledger[f"A{ledger.max_row}"].value == "total"
    figs = wb["Figures"]
    assert figs.max_row == 1 + len(ws.figures())


def test_sheet_title_is_safe_and_unique() -> None:
    taken: set[str] = set()
    a = workbook.sheet_title("tables/s1:return[levels]*?", taken)
    assert "[" not in a and ":" not in a and len(a) <= 31
    b = workbook.sheet_title("a" * 40, taken)
    c = workbook.sheet_title("a" * 40, taken)
    assert len(b) == 31 and len(c) == 31 and b != c


# ── reports ───────────────────────────────────────────────────────────────


def test_report_markdown_is_a_technical_report(ws) -> None:
    md = report_md.report_markdown(ws)
    heads = [line for line in md.splitlines() if line.startswith("## ")]
    for h in ("## Summary", "## 1 Introduction", "## 3 Data", "## 4 Method", "## 5 Results",
              "## 6 Checks and confidence", "## 8 Recommendations", "## References"):
        assert h in heads, (h, heads)
    order = [heads.index(h) for h in ("## Summary", "## 1 Introduction", "## 5 Results", "## References")]
    assert order == sorted(order)
    # the answer leads, with its unit written as a reader writes it
    assert "> **Answer**: 100-year flood: 62 m³/s" in md
    assert "m3/s" not in md and "nan" not in md
    # results are organised by question, not by step
    assert "Results: step" not in md and "### 5.1 Design flood" in md
    assert "| **100** | **1** | **62** |" in md                 # the design return period in bold
    assert "climate change" in md and "Hosking" in md
    assert "](figures/s2_frequency_curve.png)" in md and "](figures/s3_drought_strip.png)" in md
    # every figure once, numbered in order
    nums = [int(n) for n in re.findall(r"!\[Figure (\d+)\]", md)]
    assert nums == list(range(1, len(nums) + 1))


def test_report_html_is_self_contained(ws) -> None:
    html = report_md.report_html(ws)
    md = report_md.report_markdown(ws)
    assert html.count("data:image/png;base64,") == md.count("](figures/") > 0
    assert 'src="figures/' not in html and "Figure not found" not in html
    assert "<h1><span class=\"num\">5</span>Results</h1>" in html
    assert 'class="callout' in html and "@page" in html and "Hosking" in html


def test_docx_reloads_with_headings_figures_and_tables(ws) -> None:
    import docx

    data = report_docx.report_docx_bytes(ws)
    assert data is not None and data[:2] == b"PK"
    doc = docx.Document(io.BytesIO(data))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for h in ("Summary", "1  Introduction", "5  Results", "6  Checks and confidence", "References"):
        assert h in headings, (h, headings)
    assert sum(h.startswith("Appendix") for h in headings) >= 1
    assert len(doc.inline_shapes) == report_md.report_markdown(ws).count("](figures/")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Table 1." in text and "Figure 1." in text
    assert "No language model was used" in text and "m3/s" not in text
    tables_text = " ".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "Checked by" in tables_text and "Not yet checked" in tables_text
    assert doc.sections[0].different_first_page_header_footer


def test_memo_is_short_and_leads_with_the_answer(ws) -> None:
    from aquascope.studio.document import build_memo

    memo = build_memo(ws)
    assert memo.kind == "Technical memorandum"
    assert [h for h in memo.outline()][:2] == ["1 Question", "2 Answer"]
    assert len(memo.figures) <= 1 and len(memo.tables) <= 1


def test_docx_markdown_converter() -> None:
    import docx

    doc = docx.Document()
    report_docx.markdown_to_docx(doc, "Plain **bold** and `code`.\n\n- first\n- second\n\n1. one\n2. two\n\n### Sub")
    paras = doc.paragraphs
    assert paras[0].runs[1].bold and paras[0].runs[1].text == "bold"
    assert paras[0].runs[3].font.name == "Courier New"
    assert [p.style.name for p in paras[1:3]] == ["List Bullet", "List Bullet"]
    assert [p.style.name for p in paras[3:5]] == ["List Number", "List Number"]
    assert paras[5].style.name.startswith("Heading") and paras[5].text == "Sub"


def test_docx_skipped_without_python_docx(ws, monkeypatch) -> None:
    monkeypatch.setitem(sys.modules, "docx", None)
    n = len(ws.events)
    assert report_docx.report_docx_bytes(ws) is None
    assert ws.events[n]["event"] == "note" and "python-docx" in ws.events[n]["detail"]


def test_blocks_order_places_the_lists(ws) -> None:
    kinds = [(b["kind"], b.get("id") or b.get("key")) for b in _common.blocks(ws)]
    assert kinds.index(("list", "not_established")) == kinds.index(("section", "limitations")) + 1
    assert kinds.index(("list", "caveats")) == kinds.index(("section", "limitations")) + 2
    assert kinds.index(("list", "references")) == kinds.index(("section", "appendix")) - 1
    assert ("list", "recommendations") not in kinds  # the Author wrote that section


# ── notebook ──────────────────────────────────────────────────────────────


def test_notebook_is_valid_nbformat(ws) -> None:
    text = notebook.notebook_json(ws)
    nb = json.loads(text)
    assert nb["nbformat"] == 4 and nb["nbformat_minor"] == 5 and isinstance(nb["metadata"], dict)
    cells = nb["cells"]
    assert cells[0]["cell_type"] == "markdown" and "Design flow at Kingston" in cells[0]["source"]
    assert len({c["id"] for c in cells}) == len(cells)
    for c in cells:
        assert {"id", "cell_type", "source", "metadata"} <= set(c)
        if c["cell_type"] == "code":
            assert c["outputs"] == [] and c["execution_count"] is None
    code = "\n".join(c["source"] for c in cells if c["cell_type"] == "code")
    assert 'rerun_workspace("study.yaml", "workspace.json")' in code and "export(ws, out)" in code
    assert code.count("results.get(") == 3 and "plt.imread(io.BytesIO(artifact.data)" in code
    assert 'rerun-{ws.id}' in code
    nbformat = pytest.importorskip("nbformat")
    nbformat.validate(nbformat.reads(text, as_version=4))


# ── bundle ────────────────────────────────────────────────────────────────


def test_build_adds_the_documents_and_the_zip_lists_them(ws) -> None:
    ws = Workspace.from_json(ws.to_json())  # a copy: the fixture is shared by the module, and order varies
    before = len(ws.artifacts)
    added = bundle.build(ws)
    assert [a.id for a in added] == ["report-md", "report-html", "report-docx", "memo-docx", "memo-html",
                                     "workbook", "notebook", "study", "workspace", "bundle"]
    assert [a.name for a in added] == ["report.md", "report.html", "report.docx", "memo.docx", "memo.html",
                                       "workbook.xlsx", "study.ipynb", "study.yaml", "workspace.json", "bundle.zip"]
    assert len(ws.artifacts) == before + 10
    assert ws.artifact("study").data.startswith(b"# An AquaScope study (version 3)")
    ws_json = json.loads(ws.artifact("workspace").data)
    assert ws_json["id"] == ws.id and "data" not in ws_json["artifacts"][0]
    with zipfile.ZipFile(io.BytesIO(ws.artifact("bundle").data)) as zf:
        names = set(zf.namelist())
        readme = zf.read("README.txt").decode()
    assert "README.txt" in names and "bundle.zip" not in names
    for a in ws.artifacts:
        if a.kind != "bundle":
            assert a.name in names, a.id
    assert "aquascope run study.yaml" in readme and "Design flow at Kingston" in readme
    assert any(e["event"] == "artifact" and "bundle.zip" in e["detail"] for e in ws.events)
    # building again replaces by id rather than duplicating
    bundle.build(ws, formats=["md", "zip"])
    assert len(ws.artifacts) == before + 10


def test_the_bundle_carries_engineering_inputs_for_the_main_record(ws, tmp_path, monkeypatch) -> None:
    from aquascope.io import engineering as eng

    monkeypatch.setattr(eng, "hecdss_available", lambda: False)  # the same files on every platform
    data = bundle.bundle_bytes(ws)
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = set(zf.namelist())
        readme = zf.read("README.txt").decode()
        rvt = zf.read("engineering/raven/3400TH.rvt").decode()
    for tool in ("hec-hms", "hec-ras", "hec-ssp", "dss", "swmm", "modflow6", "fews", "raven"):
        assert f"engineering/{tool}/README.txt" in names, tool
    assert "engineering/hec-ssp/b17c_settings.txt" in names
    assert ":ObservationData HYDROGRAPH 1 m3/s" in rvt and "1990-01-01 00:00:00.0 1.0" in rvt
    assert "engineering/" in readme and "from UK EA 3400TH" in readme
    bundle.export(ws, tmp_path / "eng")
    assert (tmp_path / "eng" / "engineering" / "fews" / "3400TH.xml").exists()
    empty = Workspace(site=dict(SITE))
    assert bundle.engineering_files(empty) == ({}, None)


def test_build_formats_subset(ws) -> None:
    w = Workspace.from_dict(ws.to_dict())
    added = bundle.build(w, formats=["xlsx", "ipynb"])
    assert [a.id for a in added] == ["workbook", "notebook"]


def test_export_writes_every_artifact(ws, tmp_path) -> None:
    ws = Workspace.from_json(ws.to_json())  # a copy, so the shared fixture never holds the built documents
    bundle.build(ws)
    paths = bundle.export(ws, tmp_path / "out")
    assert set(paths) == {a.id for a in ws.artifacts}
    assert (tmp_path / "out" / "figures" / "s2_frequency_curve.png").read_bytes().startswith(PNG)
    csv_text = (tmp_path / "out" / "tables" / "s2_return_levels.csv").read_text(encoding="utf-8")
    assert csv_text.startswith("T,estimator,unit,estimate,lower,upper,interval_method,confidence_level")
    assert (tmp_path / "out" / "report.html").stat().st_size > 10_000
    assert (tmp_path / "out" / "bundle.zip").exists()


def test_workspace_roundtrip_keeps_artifact_bytes(ws) -> None:
    again = Workspace.from_json(ws.to_json())
    fig = again.artifact("fig-s2-series")
    assert fig is not None and fig.data == ws.artifact("fig-s2-series").data
    assert report_md.report_html(again) == report_md.report_html(ws)
