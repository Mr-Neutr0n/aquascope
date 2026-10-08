"""The documents a study becomes (aquascope.studio.document): the report, the memo, the study note.

Each case is a small hand-built workspace, so what is checked is the editorial contract: the answer first, results
by question, numbers as a reader writes them, a failed study never dressed as a report, an ungauged flood question
answered honestly with screening values and the record to get.
"""

from __future__ import annotations

import io

import pytest

from aquascope.studio import document as docs
from aquascope.studio.document import text as tx
from aquascope.studio.document.model import Callout, Table
from aquascope.studio.progress import Narrator
from aquascope.studio.workspace import Brief, Dataset, Inventory, Workspace

T = [2, 5, 10, 25, 50, 100]
AMAX_YEARS = list(range(1926, 2026))
AMAX = [200.0 + (i * 37 % 300) for i in range(100)]


def _ffa_payload() -> dict:
    return {
        "source": "usgs", "station_id": "USGS-01013500", "station_name": "Fish River near Fort Kent, Maine",
        "agency": "U.S. Geological Survey", "variable": "discharge", "unit": "m3/s",
        "start": "1903-07-29", "end": "2026-10-06", "years": 123.2, "n": 37549,
        "eligibility": {"complete_years": 100, "coverage_rule": "At least 292 observed days in each included year",
                        "scope": "Exploratory daily-mean flood screening, not regulatory design certification"},
        "annual_max": {"year": AMAX_YEARS, "v": AMAX},
        "ffa": {
            "n_years": 100, "return_periods": T,
            "fits": {
                "gev_lmoments": {"estimator": "gev_lmoments", "q": [233.4, 301.0, 343.3, 394.1, 429.9, 463.9],
                                 "params": [0.0609, 210.4844, 63.1767]},
                "lp3": {"estimator": "lp3_log_moments", "ci_level": 0.9, "interval_method": "variance_of_estimate",
                        "q": [236.3, 303.9, 343.1, 387.5, 417.6, 445.4],
                        "ci": [[224.2, 249.1], [285.6, 323.4], [319.5, 368.4], [357.0, 420.7], [381.9, 456.6],
                               [404.7, 490.2]],
                        "params": [-0.356, 2.3653, 0.1374]},
            },
            "record_max": {"value": 507.0, "year": 2008, "empirical_return_period": 101.0},
            "amax_trend": {"on": "annual maxima", "p_value": 0.3197, "tau": 0.0677, "trend": "no trend",
                           "sens_slope_per_year": 0.24, "n_years": 100},
            "amax_change": {"test": "Pettitt", "change_year": 1968, "p_value": 0.3125, "significant": False},
        },
    }


def _gauged() -> Workspace:
    ws = Workspace(site={"lat": 47.2375, "lon": -68.58278})
    ws.brief = Brief(problem="What is the 100-year flood for a culvert here?", decision="design flow",
                     quantities=["the 100-year return level with its interval"], kind="flood_risk",
                     playbook="flood_risk", intake={"return_period": 100}, ready=True)
    ws.study = {"version": 3, "title": "t", "question": "q", "plan": {"playbook": "flood_risk", "caveats": [
        "Rare quantiles move with the distribution and the estimator."]},
                "steps": [{"id": "s1", "tool": "flood_frequency", "method": "at_site_flood_frequency",
                           "arguments": {"source": "usgs", "station_id": "USGS-01013500"}}]}
    ws.run = {"ok": True, "results": [{
        "id": "s1", "tool": "flood_frequency", "ok": True, "result": _ffa_payload(),
        "gates": [{"check": "min_years", "value": 20, "passed": True, "detail": "100 years of record, 20 needed"},
                  {"check": "cross_check_ratio", "value": 0.5, "passed": True,
                   "detail": "skipped: no comparable model cell verified"}]}]}
    ws.findings = {"primary_step": "s1", "decision": {
        "answer": "design flow: 100-year return level, GEV (L-moments) 463.9 m3/s (indicative).",
        "value": 463.9, "unit": "m3/s", "grade": "indicative", "basis": ["s1.ffa.fits.gev_lmoments.q.5"],
        "conditions": ["The annual maxima are independent and drawn from a stationary distribution."],
        "limitations": [], "what_would_change_it": []}}
    ws.report = {"answer": "x", "references": ["Hosking, J. R. M. (1990). L-moments: analysis and estimation of "
                                                "distributions. J. R. Stat. Soc. B, 52(1), 105-124.",
                                                "Kendall (1975)", "a note without an author (2024), doi:10.1/x"]}
    ws.status = "done"
    return ws


def _ungauged() -> Workspace:
    ws = Workspace(site={"lat": -35.215647, "lon": 138.833176})
    ws.brief = Brief(problem="is flooding here getting better?", decision="design flow", kind="flood_risk",
                     playbook="flood_risk", intake={"return_period": 100}, ready=True)
    ws.study = {"version": 3, "plan": {"playbook": "flood_risk"},
                "steps": [{"id": "s1", "tool": "describe_catchment"}, {"id": "s2", "tool": "regionalize_signatures"}]}
    est = {"q_annual_max_mm": {"value": 4.1481, "unit": "mm/d", "label": "mean annual daily maximum",
                               "low": 2.3995, "high": 7.1712, "n_donors": 5},
           "q_mean_mm": {"value": 0.5759, "unit": "mm/d", "label": "mean daily flow", "low": 0.4041, "high": 0.8208}}
    ws.run = {"ok": True, "results": [
        {"id": "s1", "tool": "describe_catchment", "ok": True,
         "result": {"attributes": {"upstream_area_km2": 138.2, "elevation_m": {"value": 279.0, "unit": "m"}}},
         "gates": [{"check": "not_empty", "passed": True, "detail": "'sub_basin' is present", "path": "sub_basin"}]},
        {"id": "s2", "tool": "regionalize_signatures", "ok": True,
         "result": {"estimates": est, "k": 5, "skill": {"by_signature": {"q_annual_max_mm": {"nse": 0.29}}}},
         "gates": []}]}
    ws.findings = {"decision": {"answer": "No number in the results answers the decision (screening).",
                                "value": None, "grade": "screening", "basis": []}}
    ws.inventory = Inventory(site=dict(ws.site), datasets=[
        Dataset(id="south_africa_dws:A4261794:discharge", kind="station", variable="discharge",
                source="south_africa_dws",
                station_id="A4261794", name="Paris Creek Tributary upstream Paris Creek Road", distance_km=0.0)])
    ws.status = "done"
    return ws


def _failed() -> Workspace:
    ws = Workspace(site={"lat": -23.8681, "lon": -46.4011})
    ws.brief = Brief(problem="is flooding getting better", kind="flood_risk", playbook="flood_risk", ready=True)
    ws.study = {"version": 3, "plan": {}, "steps": [{"id": "s1", "tool": "describe_catchment"}]}
    ws.run = {"ok": False, "results": [{
        "id": "s1", "tool": "describe_catchment", "ok": False,
        "error": "catchment lookup failed: DataSourceError: CURL error: Could not resolve host: huggingface.co",
        "gates": [{"check": "not_empty", "passed": False, "detail": "the step returned an error"}]}]}
    ws.findings = {"decision": {"answer": "No number", "value": None, "grade": "not_established"}}
    return ws


# ── text ────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("x, want", [(463.9272, "464"), (41.21, "41.2"), (0.06361, "0.0636"), (2320.4, "2,320"),
                                     (43.0, "43"), (4.20, "4.2"), (None, "n/a")])
def test_numbers_are_written_to_three_significant_figures(x, want):
    assert tx.num(x) == want


def test_a_stated_condition_reads_as_an_assumption():
    assert tx.assumption("The annual maxima are independent and drawn from a stationary distribution.") == \
        "The estimate assumes that the annual maxima are independent and drawn from a stationary distribution."
    assert tx.assumption("Daily mean maxima can understate peaks.") == "Daily mean maxima can understate peaks."


# ── the report ──────────────────────────────────────────────────────────────


def test_the_report_leads_with_the_answer_and_is_organised_by_question():
    doc = docs.build_report(_gauged())
    assert doc.kind == "Technical report" and doc.title == "100-year design flood estimate"
    assert doc.subtitle == "Fish River near Fort Kent, Maine"
    outline = doc.outline()
    assert outline[0] == "Summary" and "5.1 Design flood" in outline and "5.2 Sensitivity of the design value" \
        in outline and "5.3 Stationarity of the annual maxima" in outline
    assert not any("step" in h.lower() for h in outline)
    box = next(b for b in doc.blocks if isinstance(b, Callout))
    assert box.body == "100-year flood: 464 m³/s"
    assert ("90 % interval", "405 to 490 m³/s (Log-Pearson III)") in box.rows
    levels = next(b for b in doc.tables if b.id == "levels")
    assert levels.rows[levels.emphasis[0]][0] == "100" and "405–490" in levels.rows[levels.emphasis[0]]


def test_the_report_text_is_written_for_a_reader():
    md = docs.render_markdown(docs.build_report(_gauged()))
    assert "m3/s" not in md and "m³/s" in md
    assert "The estimate assumes that the annual maxima are independent" in md
    assert "Mann-Kendall p = 0.32" in md and "Pettitt p = 0.31" in md
    assert "USGS 01013500" in md and "usgs USGS-01013500" not in md
    # references: the abbreviated one completed, the authorless note dropped, sorted
    assert "Kendall, M. G. (1975). Rank Correlation Methods" in md and "a note without an author" not in md
    # a skipped check is said as a limitation, in words
    assert "A check could not be run" in md


def test_cross_references_resolve_and_figures_and_tables_are_numbered():
    doc = docs.build_report(_gauged())
    md = docs.render_markdown(doc)
    assert "{tab:" not in md and "{fig:" not in md
    assert [t.number for t in doc.tables] == list(range(1, len(doc.tables) + 1))


def test_the_html_prints_and_the_word_file_reopens():
    pytest.importorskip("docx")
    import docx

    ws = _gauged()
    html = docs.render_html(docs.build_report(ws))
    assert "@page" in html and '<table class="data">' in html and 'class="callout answer"' in html
    data = docs.render_docx(docs.build_report(ws))
    word = docx.Document(io.BytesIO(data))
    heads = [p.text for p in word.paragraphs if p.style.name.startswith("Heading")]
    assert "5  Results" in heads and "Summary" in heads
    footer_xml = word.sections[0].footer._element.xml
    assert "NUMPAGES" in footer_xml and "PAGE" in footer_xml
    assert word.core_properties.title.startswith("100-year design flood estimate")


def test_the_memo_is_the_answer_on_two_pages():
    memo = docs.build_memo(_gauged())
    assert memo.kind == "Technical memorandum"
    assert memo.outline()[:2] == ["1 Question", "2 Answer"]
    levels = next(b for b in memo.blocks if isinstance(b, Table))
    assert len(levels.rows) <= 4 and levels.rows[levels.emphasis[0]][0] == "100"


# ── an ungauged flood question ──────────────────────────────────────────────


def test_an_ungauged_flood_question_is_answered_honestly():
    ws = _ungauged()
    doc = docs.build_report(ws)
    assert doc.title == "Flood screening at an ungauged site"
    box = next(b for b in doc.blocks if isinstance(b, Callout))
    assert box.body.startswith("No design flood") and box.tone == "caution"
    rows = dict(box.rows)
    # 4.1481 mm/d over 138.2 km2 is 6.64 m3/s (Q = depth x area / 86.4)
    assert rows["Screening value"].startswith("mean annual maximum daily flow 6.64 m³/s")
    assert "Paris Creek Tributary" in rows["Next step"] and "at the site" in rows["Next step"]
    md = docs.render_markdown(doc)
    assert "No number in the results" not in md
    assert "Ask the agency for its daily record" in md       # AquaScope cannot fetch this source yet
    assert "Do not size a structure from the screening values" in md


# ── a study that established nothing ────────────────────────────────────────


def test_a_failed_study_becomes_a_note_not_a_report():
    ws = _failed()
    assert docs.is_failed_study(ws)
    doc = docs.build_report(ws)
    assert doc.kind == "Study note" and doc.status == "NOT ESTABLISHED"
    md = docs.render_markdown(doc)
    assert "CURL" not in md and "could not be reached (no network connection)" in md
    assert "Check the internet connection" in md
    assert docs.build_memo(ws).kind == "Study note"


# ── house style ─────────────────────────────────────────────────────────────


def test_the_house_style_travels_with_the_workspace():
    logo = b"\x89PNG\r\n\x1a\n" + b"0" * 16
    style = docs.HouseStyle(organisation="Northern Rivers Water Authority", prepared_by="A. Hydrologist",
                         status="FINAL", logo=logo, accent="#123456")
    again = docs.HouseStyle.from_dict(style.to_dict())
    assert again.logo == logo and again.organisation == style.organisation
    html = docs.render_html(docs.build_report(_gauged(), again), again)
    assert "Northern Rivers Water Authority" in html and 'class="stamp final"' in html
    assert "A. Hydrologist" in html


def test_a_bad_accent_falls_back():
    assert docs.HouseStyle.from_dict({"accent": "blue"}).accent == docs.HouseStyle().accent


def test_the_terminal_summary_names_the_answer_and_the_files():
    ws = _gauged()
    lines = docs.terminal_summary(ws, "out")
    text = "\n".join(lines)
    assert "100-year flood: 464 m³/s" in text and "Grade:" in text and "Documents in out:" in text


# ── the progress log ────────────────────────────────────────────────────────


def test_the_narrator_folds_gates_into_one_line_per_step():
    n = Narrator()
    lines: list[str] = []
    for e in [
        {"role": "consultant", "event": "brief", "detail": "flood risk"},
        {"role": "scout", "event": "inventory", "detail": "9 dataset(s): a, b"},
        {"role": "runner", "step": "s1", "event": "start", "detail": "flood_frequency(source='usgs')"},
        {"role": "runner", "step": "s1", "event": "done", "detail": "ok"},
        {"role": "reviewer", "step": "s1", "event": "gate", "detail": "min_years: passed, 100 years of record"},
        {"role": "reviewer", "step": "s1", "event": "gate", "detail": "cross_check_ratio: passed, skipped: no cell"},
        {"role": "reviewer", "step": "s1", "event": "gate", "detail": "spread_within: failed, spread 40%"},
        {"role": "analyst", "event": "gates", "detail": "1 of 3"},
        {"role": "author", "event": "artifact", "detail": "report.docx (1 bytes)"},
    ]:
        lines += n.feed(e)
    lines += n.flush()
    assert lines[0] == "Reading the question" and lines[1].startswith("Looking for data near the site: 9")
    step = next(x for x in lines if "Flood frequency analysis" in x)
    assert step.startswith("✗") and "1 passed, 1 failed, 1 could not be run" in step
    assert any(x.strip().startswith("✗ spread_within") for x in lines)
    assert not any("report.docx" in x for x in lines)
