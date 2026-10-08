"""A document as a Word file a hydrologist can edit, sign and issue.

Real Word structure, not pasted text: the body in the Normal style, headings
in Heading 1/2 (so the navigation pane and a table of contents work), tables
as Word tables in the booktabs pattern (a heavy rule above and below, a thin
one under the header, nothing else), captions above tables and below
figures, a cover with the document-control table and sign-off lines, a
running header with the organisation and the title, and a footer with the
status and "Page X of Y" as live fields. Figures are inserted at the text
width at 300 dpi. python-docx is imported here only.
"""

from __future__ import annotations

import io
import re
from typing import Any

from aquascope.studio.document.model import (
    Bullets,
    Callout,
    Document,
    Figure,
    Heading,
    KeyValue,
    PageBreak,
    Para,
    Signoff,
    Table,
)
from aquascope.studio.document.style import HouseStyle

__all__ = ["render_docx"]

INK = (0x1A, 0x1A, 0x1A)
MUTED = (0x5B, 0x67, 0x70)
FAIL = (0xB3, 0x26, 0x1E)
CAUTION = (0xB2, 0x6A, 0x00)

_INLINE = re.compile(r"(\*\*.+?\*\*|(?<![*\w])\*(?!\s).+?(?<!\s)\*(?![*\w]))")


def _rgb(hex_or_tuple: Any) -> Any:
    from docx.shared import RGBColor

    if isinstance(hex_or_tuple, tuple):
        return RGBColor(*hex_or_tuple)
    h = str(hex_or_tuple).lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_font(run_or_style: Any, name: str) -> None:
    """Set a font on every script slot, so Word never falls back to Calibri on symbols or East Asian text."""
    from docx.oxml.ns import qn

    run_or_style.font.name = name
    el = run_or_style.element if hasattr(run_or_style, "element") else run_or_style._element
    rpr = el.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        from docx.oxml import OxmlElement

        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(slot), name)


def _runs(par: Any, text: str, *, size: float | None = None, color: Any = None, font: str | None = None,
          bold: bool | None = None) -> None:
    from docx.shared import Pt

    for part in _INLINE.split(str(text or "")):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run, b, it = par.add_run(part[2:-2]), True, False
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run, b, it = par.add_run(part[1:-1]), False, True
        else:
            run, b, it = par.add_run(part), False, False
        if b or bold:
            run.bold = True
        if it:
            run.italic = True
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = _rgb(color)
        if font:
            _set_font(run, font)


def _border(cell: Any, **edges: tuple[int, str]) -> None:
    """Set cell borders: ``top=(size_eighths, "000000")``; size 0 clears the edge."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge, (size, color) in edges.items():
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        if size <= 0:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), color)
            el.set(qn("w:space"), "0")


def _shade(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill.lstrip("#"))
    tcpr.append(shd)


def _no_table_borders(table: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tblpr = table._tbl.tblPr
    for old in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(old)
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        b.append(el)
    tblpr.append(b)


def _cell_margins(table: Any, top: int = 40, bottom: int = 40, left: int = 80, right: int = 80) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def _field(par: Any, code: str, *, size: float, font: str, color: Any) -> None:
    """A live Word field (PAGE, NUMPAGES) in a paragraph."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    def run_with(el: Any) -> None:
        r = par.add_run()
        r.font.size = Pt(size)
        r.font.color.rgb = _rgb(color)
        _set_font(r, font)
        r._r.append(el)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {code} "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, txt, end):
        run_with(el)


def _keep_with_next(par: Any) -> None:
    par.paragraph_format.keep_with_next = True


def _styles(doc: Any, style: HouseStyle) -> None:
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.shared import Pt

    normal = doc.styles["Normal"]
    _set_font(normal, style.body_font)
    normal.font.size = Pt(style.body_size)
    normal.font.color.rgb = _rgb(INK)
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.18
    for name, size, before, after in (("Heading 1", 13.5, 16, 6), ("Heading 2", 11.5, 12, 4),
                                      ("Heading 3", 10.5, 10, 3), ("Title", 22, 0, 4), ("Subtitle", 13, 0, 10)):
        st = doc.styles[name]
        _set_font(st, style.heading_font)
        st.font.size = Pt(size)
        st.font.bold = name != "Subtitle"
        st.font.italic = False
        st.font.color.rgb = _rgb(MUTED if name == "Subtitle" else INK)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        for border in list(st.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    for name in ("List Bullet", "List Number"):
        st = doc.styles[name]
        st.paragraph_format.space_after = Pt(3)


def _page(doc: Any, style: HouseStyle, title: str, status: str) -> None:
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Cm, Mm, Pt

    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    if style.paper.upper() == "A4":
        sec.page_width, sec.page_height = Mm(210), Mm(297)
    else:
        sec.page_width, sec.page_height = Mm(215.9), Mm(279.4)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.3)
    sec.header_distance = sec.footer_distance = Cm(1.1)
    sec.different_first_page_header_footer = True
    width = sec.page_width - sec.left_margin - sec.right_margin

    head = sec.header.paragraphs[0]
    head.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    _runs(head, style.author_line, size=7.5, color=MUTED, font=style.heading_font)
    _runs(head, "\t" + title, size=7.5, color=MUTED, font=style.heading_font)
    for footer in (sec.footer, sec.first_page_footer):
        foot = footer.paragraphs[0]
        foot.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        if status:
            r = foot.add_run(status.upper())
            r.bold = True
            r.font.size = Pt(7.5)
            r.font.color.rgb = _rgb(FAIL if status.upper() not in ("FINAL", "ISSUED") else style.accent)
            _set_font(r, style.heading_font)
        _runs(foot, "\tPage ", size=7.5, color=MUTED, font=style.heading_font)
        _field(foot, "PAGE", size=7.5, font=style.heading_font, color=MUTED)
        _runs(foot, " of ", size=7.5, color=MUTED, font=style.heading_font)
        _field(foot, "NUMPAGES", size=7.5, font=style.heading_font, color=MUTED)


def _cover(doc: Any, d: Document, style: HouseStyle) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    if style.logo:
        try:
            doc.add_picture(io.BytesIO(style.logo), height=Inches(0.55))
        except Exception:  # noqa: BLE001 - an unreadable logo is left out, not fatal
            pass
    org = doc.add_paragraph()
    _runs(org, style.author_line, size=9, color=MUTED, font=style.heading_font)
    if d.status:
        stamp = doc.add_paragraph()
        stamp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _runs(stamp, d.status.upper(), size=9, bold=True, font=style.heading_font,
              color=style.accent if d.status.upper() in ("FINAL", "ISSUED") else
              CAUTION if d.status.upper() == "CHECKED" else FAIL)
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(36)
    kicker.paragraph_format.space_after = Pt(4)
    _runs(kicker, d.kind.upper(), size=8.5, bold=True, color=style.accent, font=style.heading_font)
    t = doc.add_paragraph(style="Title")
    _runs(t, d.title)
    if d.subtitle:
        s = doc.add_paragraph(style="Subtitle")
        _runs(s, d.subtitle)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(14)
    _para_rule(rule, "bottom", 12)


def _para_rule(par: Any, edge: str, size: int, color: str = "1A1A1A") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ppr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "1")
    el.set(qn("w:color"), color)
    bdr.append(el)
    ppr.append(bdr)


def _caption(doc: Any, label: str, text: str, style: HouseStyle, *, above: bool) -> None:
    from docx.shared import Pt

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8 if above else 3)
    p.paragraph_format.space_after = Pt(4 if above else 12)
    if above:
        _keep_with_next(p)
    if label:
        _runs(p, label + "  ", size=style.body_size - 1.5, bold=True, font=style.heading_font)
    _runs(p, text, size=style.body_size - 1.5)


def _table(doc: Any, t: Table, style: HouseStyle) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    _caption(doc, f"Table {t.number}." if t.numbered and t.number else "", t.caption, style, above=True)
    n_cols = len(t.columns)
    table = doc.add_table(rows=1 + len(t.rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER if t.compact else WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    _no_table_borders(table)
    _cell_margins(table)
    size = style.body_size - (2.5 if n_cols >= 7 else 2.0)
    aligns = {"r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}
    for r_i, values in enumerate([t.columns, *t.rows]):
        for c_i in range(n_cols):
            cell = table.cell(r_i, c_i)
            v = values[c_i] if c_i < len(values) else ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            a = t.align[c_i] if c_i < len(t.align) else "l"
            if a in aligns:
                par.alignment = aligns[a]
            bold = r_i == 0 or (r_i - 1) in t.emphasis
            _runs(par, "" if v is None else str(v), size=size, bold=bold, font=style.heading_font)
            edges: dict[str, tuple[int, str]] = {}
            if r_i == 0:
                edges["top"] = (12, "1A1A1A")
                edges["bottom"] = (6, "1A1A1A")
            if r_i == len(t.rows):
                edges["bottom"] = (12, "1A1A1A")
            if edges:
                _border(cell, **edges)
        if r_i == 0:
            _repeat_header(table.rows[0])
    for note in t.notes:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        _runs(p, note, size=style.body_size - 2.5, color=MUTED)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def _repeat_header(row: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _figure(doc: Any, f: Figure, style: HouseStyle, text_width: Any) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    width = text_width if f.width == "full" else int(text_width * 0.5)
    doc.add_picture(io.BytesIO(f.png), width=width)
    pic = doc.paragraphs[-1]
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_before = Pt(6)
    pic.paragraph_format.space_after = Pt(0)
    _keep_with_next(pic)
    _caption(doc, f"Figure {f.number}.", f.caption, style, above=False)


def _callout(doc: Any, c: Callout, style: HouseStyle) -> None:
    from docx.shared import Pt

    colour = {"answer": style.accent, "caution": "#B26A00", "fail": "#B3261E"}.get(c.tone, style.accent)
    fill = {"answer": "#F3F6F8", "caution": "#FBF5EC", "fail": "#FBEFEE"}.get(c.tone, "#F3F6F8")
    table = doc.add_table(rows=1, cols=1)
    _no_table_borders(table)
    _cell_margins(table, top=120, bottom=100, left=200, right=160)
    cell = table.cell(0, 0)
    _shade(cell, fill)
    _border(cell, left=(36, colour.lstrip("#")))
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    _runs(p, c.title.upper(), size=8, bold=True, color=colour, font=style.heading_font)
    if c.body:
        b = cell.add_paragraph()
        b.paragraph_format.space_after = Pt(5)
        _runs(b, c.body, size=14, bold=True, font=style.heading_font)
    if c.rows:
        from docx.shared import Cm

        inner = cell.add_table(rows=len(c.rows), cols=2)
        _no_table_borders(inner)
        _cell_margins(inner, top=10, bottom=10, left=0, right=60)
        _fixed_widths(inner, [Cm(3.4), Cm(11.4)])
        for i, (label, value) in enumerate(c.rows):
            for col, text, colour, bold in ((0, label, MUTED, True), (1, value, INK, False)):
                par = inner.cell(i, col).paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                _runs(par, text, size=style.body_size - 2, bold=bold, color=colour, font=style.heading_font)
        cell.add_paragraph().paragraph_format.space_after = Pt(0)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(4)


def _fixed_widths(table: Any, widths: list[Any]) -> None:
    """Fixed column widths that every renderer honours (Word, LibreOffice, Quick Look): the grid, each cell, and
    a fixed layout."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table.autofit = False
    tblpr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    grid = table._tbl.tblGrid
    for col, w in zip(grid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(w.twips if hasattr(w, "twips") else w / 635)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def _keyvalue(doc: Any, kv: KeyValue, style: HouseStyle) -> None:
    from docx.shared import Cm, Pt

    table = doc.add_table(rows=len(kv.rows), cols=2)
    _no_table_borders(table)
    _cell_margins(table, top=30, bottom=30, left=0, right=80)
    _fixed_widths(table, [Cm(4.0), Cm(12.0)])
    for i, (k, v) in enumerate(kv.rows):
        a, b = table.cell(i, 0), table.cell(i, 1)
        for cell, text, colour, bold in ((a, k, MUTED, True), (b, v, INK, False)):
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            _runs(par, str(text), size=9, color=colour, bold=bold, font=style.heading_font)
            _border(cell, bottom=(2, "D5DADE"))
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(6)


def _signoff(doc: Any, s: Signoff, style: HouseStyle) -> None:
    from docx.shared import Cm, Pt

    table = doc.add_table(rows=3, cols=len(s.roles))
    _no_table_borders(table)
    _cell_margins(table, top=40, bottom=20, left=0, right=200)
    _fixed_widths(table, [Cm(16.0 / len(s.roles))] * len(s.roles))
    for c, (role, name) in enumerate(s.roles):
        head = table.cell(0, c)
        _border(head, top=(6, "1A1A1A"))
        _runs(head.paragraphs[0], role, size=8.5, bold=True, font=style.heading_font)
        _runs(table.cell(1, c).paragraphs[0], name or " ", size=8.5, font=style.heading_font)
        table.cell(1, c).paragraphs[0].paragraph_format.space_after = Pt(16)
        _runs(table.cell(2, c).paragraphs[0], "Signature and date", size=7.5, color=MUTED, font=style.heading_font)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(8)


def render_docx(d: Document, style: HouseStyle | None = None) -> bytes | None:
    """The document as .docx bytes, or None when python-docx is not installed."""
    try:
        from docx import Document as WordDocument
    except ImportError:
        return None
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt

    style = style or HouseStyle()
    doc = WordDocument()
    _styles(doc, style)
    _page(doc, style, d.title + (f": {d.subtitle}" if d.subtitle else ""), d.status)
    sec = doc.sections[0]
    text_width = sec.page_width - sec.left_margin - sec.right_margin
    core = doc.core_properties
    core.title = d.title + (f": {d.subtitle}" if d.subtitle else "")
    core.subject = d.kind
    core.author = style.prepared_by or style.organisation or "AquaScope Studio"
    core.keywords = "hydrology; AquaScope"
    _cover(doc, d, style)
    for b in d.blocks:
        if isinstance(b, Heading):
            h = doc.add_heading(level=min(max(b.level, 1), 3))
            if b.number:
                _runs(h, b.number + "  ", color=style.accent)
            _runs(h, b.text)
        elif isinstance(b, Para):
            p = doc.add_paragraph()
            if b.style == "small":
                _runs(p, b.text, size=style.body_size - 2, color=MUTED)
            elif b.style == "lead":
                _runs(p, b.text, size=style.body_size + 0.5)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            else:
                _runs(p, b.text)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif isinstance(b, Bullets):
            for item in b.items:
                p = doc.add_paragraph(style="List Number" if b.numbered else "List Bullet")
                _runs(p, item)
            _restart_numbering(doc, b)
        elif isinstance(b, Figure):
            _figure(doc, b, style, text_width)
        elif isinstance(b, Table):
            _table(doc, b, style)
        elif isinstance(b, Callout):
            _callout(doc, b, style)
        elif isinstance(b, KeyValue):
            _keyvalue(doc, b, style)
        elif isinstance(b, Signoff):
            _signoff(doc, b, style)
        elif isinstance(b, PageBreak):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    last = doc.add_paragraph()
    last.paragraph_format.space_after = Pt(0)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_numbering_state: dict[int, int] = {}


def _restart_numbering(doc: Any, b: Bullets) -> None:
    """Each numbered list starts at 1: Word's List Number style continues across the document otherwise."""
    if not b.numbered:
        return
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    numbering = doc.part.numbering_part.numbering_definitions._numbering
    style_el = doc.styles["List Number"].element
    num_pr = style_el.find(qn("w:pPr") + "/" + qn("w:numPr"))
    if num_pr is None:
        return
    num_id_el = num_pr.find(qn("w:numId"))
    if num_id_el is None:
        return
    base_num = numbering.find(f"{qn('w:num')}[@{qn('w:numId')}='{num_id_el.get(qn('w:val'))}']")
    if base_num is None:
        return
    abstract = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    an = OxmlElement("w:abstractNumId")
    an.set(qn("w:val"), abstract)
    num.append(an)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    items = doc.paragraphs[-len(b.items):]
    for p in items:
        ppr = p._p.get_or_add_pPr()
        np_ = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl")
        il.set(qn("w:val"), "0")
        ni = OxmlElement("w:numId")
        ni.set(qn("w:val"), str(new_id))
        np_.append(il)
        np_.append(ni)
        ppr.append(np_)
